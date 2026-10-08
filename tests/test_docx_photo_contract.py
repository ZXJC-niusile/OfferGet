"""DOCX portrait containment and source-text contracts; actual rendering is separate."""
import base64
import copy
import io
import json
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from core.ats_docx_renderer import AtsDocxRenderer, validate_docx_ats_compliance
from core.resume_builder import ResumeBuilder
from core.resume_photo import image_to_base64_data_uri, photo_data_uri_to_png
from core.resume_view_model import ResumeViewModel
from presentation.commands.resume_cmd import cmd_resume_build
from tests.test_ats_docx_design import synthetic_profile
from tests.test_resume_photo_contract import portrait  # one shared synthetic illustration


def text_runs(doc):
    return [node.text for node in doc._element.body.iter(qn("w:t"))]


def test_photo_masthead_preserves_reading_order_and_linear_body(portrait, tmp_path):
    source = synthetic_profile()
    source["photo_path"] = str(portrait)
    original = copy.deepcopy(source)
    plain = Document(AtsDocxRenderer().render(source, tmp_path / "none.docx", photo_mode="no-photo"))
    photo_path = AtsDocxRenderer().render(source, tmp_path / "photo.docx", photo_mode="photo")
    photo = Document(photo_path)
    assert text_runs(photo) == text_runs(plain)
    assert source == original
    assert len(photo.tables) == 1 and not plain.tables
    assert len(photo.tables[0].rows) == 1 and len(photo.tables[0].columns) == 2
    assert photo._element.body[0].tag == qn("w:tbl")
    assert photo.tables[0].cell(0, 0).paragraphs[0].text == source["name"]
    assert not photo.tables[0].cell(0, 1).text
    assert len(photo.inline_shapes) == 1 and not plain.inline_shapes
    assert all(child.tag != qn("w:tbl") for child in list(photo._element.body)[1:])
    xml = photo._element.xml
    assert "w:txbxContent" not in xml and "wp:anchor" not in xml
    assert "w:cantSplit" in xml and 'descr="Candidate photo"' in xml
    assert len(list(photo._element.body.iter(qn("w:t")))) == len(list(plain._element.body.iter(qn("w:t"))))
    with zipfile.ZipFile(photo_path) as package:
        media = [name for name in package.namelist() if name.startswith("word/media/")]
        assert len(media) == 1 and media[0].endswith(".png")
        pixels = Image.open(io.BytesIO(package.read(media[0])))
        assert pixels.size == (120, 160) and pixels.getpixel((3, 3)) == (194, 82, 82)
        relationships = package.read("word/_rels/document.xml.rels").decode()
        assert "TargetMode=\"External\"" not in relationships
    # Masthead table is disclosed by the existing ATS checker; it is not an ATS guarantee.
    check = validate_docx_ats_compliance(photo_path)
    assert check["passed"] and any("table" in warning for warning in check["warnings"])


@pytest.mark.parametrize("size", [(120, 160), (400, 100), (100, 400), (100, 100)])
def test_photo_is_contained_without_crop_or_stretch(tmp_path, size):
    file = tmp_path / "ratio.png"; Image.new("RGB", size, "#416F8F").save(file)
    path = ResumeBuilder({"name": "Synthetic", "photo_path": str(file)}, photo_mode="photo").build_theme(
        "ats-docx", tmp_path / "docx", diff_ledger=False)["docx"]
    doc = Document(path); shape = doc.inline_shapes[0]
    assert shape.width.pt <= 54.01 and shape.height.pt <= 72.01
    assert shape.width / shape.height == pytest.approx(size[0] / size[1], rel=1e-4)
    assert min(54 - shape.width.pt, 72 - shape.height.pt) < 0.01
    assert "a:srcRect" not in doc._element.xml


def test_no_photo_has_no_media_no_blank_slot_and_does_not_read(portrait, tmp_path):
    source = synthetic_profile(); source["photo_path"] = str(portrait)
    clean = Document(AtsDocxRenderer().render(synthetic_profile(), tmp_path / "baseline.docx"))
    with patch("core.resume_photo.image_to_base64_data_uri", side_effect=AssertionError("must not open")):
        result = ResumeBuilder(source, photo_path=portrait, photo_mode="photo").build_theme(
            "ats-docx", tmp_path / "none", photo=False, diff_ledger=False)
    none = Document(result["docx"])
    assert text_runs(none) == text_runs(clean)
    assert [p.text for p in none.paragraphs] == [p.text for p in clean.paragraphs]
    assert not none.tables and not none.inline_shapes
    assert "预留" not in none._element.xml and "placeholder" not in none._element.xml
    with zipfile.ZipFile(result["docx"]) as package:
        assert not any(name.startswith("word/media/") for name in package.namelist())
        assert "/image" not in package.read("word/_rels/document.xml.rels").decode()


def test_validated_photo_snapshot_works_after_original_removed(portrait, tmp_path):
    vm = ResumeViewModel.from_profile({"name": "Synthetic"}, photo_path=portrait, photo_mode="photo")
    portrait.unlink()
    result = AtsDocxRenderer().render(vm, tmp_path / "snapshot.docx")
    assert len(Document(result).inline_shapes) == 1


@pytest.mark.parametrize("data", ["data:image/png;base64,synthetic", "data:text/plain;base64,aGVsbG8=", "https://example.test/photo.png"])
def test_unvalidated_inline_data_is_rejected(data, tmp_path):
    vm = ResumeViewModel(name="Synthetic", raw_name="Synthetic", reserve_photo=True,
                         photo_mode="photo", photo_data_uri=data)
    with pytest.raises(ValueError):
        AtsDocxRenderer().render(vm, tmp_path / "bad.docx")
    assert not (tmp_path / "bad.docx").exists()


def test_missing_real_photo_fails_and_placeholder_is_preview_only(tmp_path):
    with pytest.raises(ValueError, match="requires an actual"):
        AtsDocxRenderer().render({"name": "Synthetic"}, tmp_path / "missing.docx", photo_mode="photo")
    path = AtsDocxRenderer().render({"name": "Synthetic"}, tmp_path / "preview.docx", photo_mode="placeholder")
    doc = Document(path)
    assert "preview only" in " ".join(p.text for p in doc.paragraphs)
    assert not doc.inline_shapes and not doc.tables


@pytest.mark.parametrize("format,suffix", [("JPEG", ".jpg"), ("WEBP", ".webp"), ("GIF", ".gif")])
def test_supported_still_images_embed_as_portable_png(tmp_path, format, suffix):
    image = tmp_path / ("portrait" + suffix)
    Image.new("RGB", (60, 80), "#416F8F").save(image, format=format)
    path = AtsDocxRenderer().render({"name": "Synthetic", "photo_path": str(image)}, tmp_path / "photo.docx")
    with zipfile.ZipFile(path) as package:
        media = next(name for name in package.namelist() if name.startswith("word/media/"))
        assert media.endswith(".png")
        assert Image.open(io.BytesIO(package.read(media))).size == (60, 80)


def test_exif_orientation_is_normalized_before_containment(tmp_path):
    image = Image.new("RGB", (80, 60), "#416F8F")
    exif = Image.Exif(); exif[274] = 6
    file = tmp_path / "rotated.jpg"; image.save(file, exif=exif)
    data, width, height = photo_data_uri_to_png(image_to_base64_data_uri(file))
    assert (width, height) == (60, 80)
    assert Image.open(io.BytesIO(data)).size == (60, 80)


def test_mime_mismatch_rejected_even_for_valid_image_bytes(portrait):
    false_type = "data:image/jpeg;base64," + base64.b64encode(portrait.read_bytes()).decode()
    with pytest.raises(ValueError, match="MIME type"):
        photo_data_uri_to_png(false_type)


def test_cli_builds_actual_docx_photo_pair(portrait, tmp_path, capsys):
    source = synthetic_profile(); source["photo_path"] = str(portrait)
    profile_path = tmp_path / "profile.json"; profile_path.write_text(json.dumps(source))
    for mode in ("photo", "no-photo"):
        out = tmp_path / mode
        cmd_resume_build([str(profile_path), "--theme", "ats-docx", "--photo-mode", mode,
                          "--output-dir", str(out), "--no-diff-ledger"])
        assert len(Document(out / "resume_ats.docx").inline_shapes) == (1 if mode == "photo" else 0)
    assert "Successfully built" in capsys.readouterr().out
