"""DOCX content/layout contract using fabricated resume facts only."""
from pathlib import Path

import pytest

pytest.importorskip("docx")
from docx import Document
from docx.oxml.ns import qn

from core.ats_docx_renderer import AtsDocxRenderer
from core.resume_view_model import ResumeViewModel


def synthetic_profile():
    return {
        "name": "林示例 Alex & Lin", "intent": "R&D Engineer",
        "email": "qa@example.invalid", "phone": "+00 000 000 0000", "city": "示例城市",
        "education": {"school": "Example & 示例大学", "degree": "MSc", "major": "Computing",
                      "time": "2018 - 2020", "gpa": "3.8/4.0", "awards": "Test & Scholarship"},
        "experience": [{"company": "ACME & R&D", "dept": "Tools <Platform>", "role": "Engineer",
                        "time": "2020 - 2024", "bullets": [
                            "Maintained **C++ & Python**; latency < 20 ms.",
                            "Literal &amp; and <strong>source</strong> remain source text."]}],
        "campus_experience": [{"name": "Project <One>", "role": "Owner & Maintainer",
                               "time": "2019 - 2020", "bullets": ["TEST-PROJECT-FACT"]}],
        "skills": {"Tools & Languages": ["C++", "R&D"]},
        "publications": [{"title": "Safety & Systems", "authors": "A. Lin & B. Example",
                          "venue": "Test <Journal>", "year": "2025"}],
        "postgraduate_exam": {"target_school": "Example Graduate School", "target_major": "计算机",
                              "direction": "Systems & Reliability", "scores": {"Math <A>": 130, "English": 80},
                              "total_score": 210},
        "narrative": {"summary_en": "Builds reproducible systems for synthetic QA."},
    }


def render(tmp_path, profile=None):
    path = AtsDocxRenderer().render(profile or synthetic_profile(), tmp_path / "resume.docx")
    return Document(path)


def test_docx_preserves_raw_source_facts(tmp_path):
    doc = render(tmp_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for expected in (
        "ACME & R&D", "Tools <Platform>", "Engineer", "2020 - 2024", "C++ & Python",
        "latency < 20 ms.", "Literal &amp; and <strong>source</strong>",
        "Project <One>", "Owner & Maintainer", "2019 - 2020", "TEST-PROJECT-FACT",
        "Example & 示例大学", "MSc", "Computing", "2018 - 2020", "3.8/4.0", "Test & Scholarship",
        "Safety & Systems", "A. Lin & B. Example", "Test <Journal>", "2025",
        "Example Graduate School", "计算机", "Systems & Reliability", "Math <A>: 130", "English: 80", "210",
        "Tools & Languages: C++, R&D", "Builds reproducible systems",
    ):
        assert expected in text, expected
    assert "&lt;" not in text
    assert "<strong>C++" not in text
    assert "**C++" not in text
    assert any(r.bold and r.text == "C++ & Python" for p in doc.paragraphs for r in p.runs)


def test_legacy_html_projected_view_model_is_decoded_once(tmp_path):
    vm = ResumeViewModel(name="A &amp; B", experiences=[{
        "org": "ACME &amp; Co", "role": "R&amp;D", "time": "2020 &lt; 2025",
        "bullets": ["Used <strong>C++ &amp; SQL</strong> with &lt; 1 error."]}])
    doc = render(tmp_path, vm)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ACME & Co" in text and "2020 < 2025" in text
    assert "Used C++ & SQL with < 1 error." in text
    assert any(r.bold and r.text == "C++ & SQL" for p in doc.paragraphs for r in p.runs)


def test_docx_uses_a4_linear_flow_and_safe_pagination(tmp_path):
    doc = render(tmp_path)
    section = doc.sections[0]
    assert section.page_width.mm == pytest.approx(210, abs=0.1)
    assert section.page_height.mm == pytest.approx(297, abs=0.1)
    assert section.left_margin.inches == pytest.approx(0.68, abs=0.001)
    assert not doc.tables and not doc.inline_shapes
    assert not section.header.paragraphs[0].text
    assert not section.footer.paragraphs[0].text
    headings = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings and all(p.paragraph_format.keep_with_next for p in headings)
    assert all(p._p.xpath("./w:pPr/w:pBdr/w:bottom") for p in headings)
    assert not doc.styles["Title"].element.xpath("./w:pPr/w:pBdr")
    dates = next(p for p in doc.paragraphs if "ACME & R&D" in p.text)
    assert dates.text == "ACME & R&D\t2020 - 2024"
    assert dates.paragraph_format.keep_with_next
    assert dates._p.xpath("./w:pPr/w:tabs/w:tab")[0].get(qn("w:val")) == "right"
    bullet = next(p for p in doc.paragraphs if "latency" in p.text)
    assert not bullet.paragraph_format.keep_with_next
    assert bullet.paragraph_format.keep_together
    assert bullet.style.paragraph_format.first_line_indent < 0


def test_photo_preferences_do_not_introduce_ats_text_boxes(tmp_path):
    vm = ResumeViewModel.from_profile(synthetic_profile(), photo_mode="no-photo")
    vm.photo_data_uri = "data:image/png;base64,synthetic"
    doc = render(tmp_path, vm)
    assert not doc.inline_shapes and not doc.tables
    assert "ACME & R&D" in "\n".join(p.text for p in doc.paragraphs)


def test_long_item_labels_get_a_separate_unfragmented_date_line(tmp_path):
    profile = synthetic_profile()
    profile["experience"][0]["company"] = "LongInstitutionName" * 10
    doc = render(tmp_path, profile)
    paragraph = next(p for p in doc.paragraphs if p.text == "2020 - 2024")
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraph.paragraph_format.keep_with_next


@pytest.mark.parametrize("source", ["中文。 English text", "**中文。 English text**",
                                    "中文。 **English text**", "**中文。** English text"])
def test_bilingual_break_preserves_text_and_emphasis(tmp_path, source):
    profile = synthetic_profile()
    profile["experience"][0]["bullets"] = [source, "Used C++ and Python. Kept original terms."]
    doc = render(tmp_path, profile)
    bullet = next(p for p in doc.paragraphs if "中文。" in p.text)
    assert bullet.text == "中文。\nEnglish text"
    english = next(r for r in bullet.runs if r.text == "English text")
    assert bool(english.bold) == ("text**" in source)
    assert english.font.size.pt == 10
    assert any(p.text == "Used C++ and Python. Kept original terms." for p in doc.paragraphs)


def test_sparse_identity_and_publications_keep_semantic_structure(tmp_path):
    vm = ResumeViewModel(raw_email="qa@example.invalid", publications=[
        {"title": "Literal · separator", "status": "Submitted, not accepted",
         "individual_contribution": "Reviewed · outputs"}, {"author_role": "Supporting author"}])
    doc = render(tmp_path, vm)
    assert doc.paragraphs[0].text == "qa@example.invalid"
    assert doc.paragraphs[0].style.name != "Title"
    entry = next(p for p in doc.paragraphs if p.text.startswith("Literal · separator"))
    assert entry.text == "Literal · separator\n状态：Submitted, not accepted\n个人贡献：Reviewed · outputs"
    assert entry.runs[0].bold and entry.paragraph_format.keep_together
    assert not entry.paragraph_format.keep_with_next
    assert any(p.text == "作者角色：Supporting author" for p in doc.paragraphs)
    oversized = ResumeViewModel(publications=[{"title": "Long entry", "summary": "Long summary 中文。 " * 500}])
    long_doc = render(tmp_path, oversized)
    long_entry = next(p for p in long_doc.paragraphs if p.text.startswith("Long entry"))
    assert not long_entry.paragraph_format.keep_together
    assert long_entry.text.count("Long summary 中文。") == 500
